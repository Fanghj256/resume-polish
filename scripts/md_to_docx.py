"""Markdown → Word（.docx）：给衣柜、打分表、质检这类文档导出一份可读的 Word 版。

真相源永远是 markdown；本脚本只做单向导出，不要反向手工编辑 docx。
支持：# 标题、| 表格 |、- 列表、> 引用、**加粗**、`代码`、[文字](链接)。

用法：
    python scripts/md_to_docx.py --md="/path/to/简历衣柜.md" --out="/path/to/简历衣柜.docx"

依赖：python-docx（python -m pip install python-docx）
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

CJK_FONT = "微软雅黑"  # macOS 可改成 "PingFang SC"；缺失时 Word 会自动替换
BODY_PT = 10.5
TABLE_PT = 9


def set_cjk_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = CJK_FONT
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)


def add_runs(paragraph, text: str, size: float | None = None) -> None:
    """写入一行文字，处理 **加粗**、`代码`、[文字](链接)。"""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for chunk in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        bold = chunk.startswith("**") and chunk.endswith("**")
        code = chunk.startswith("`") and chunk.endswith("`")
        run = paragraph.add_run(chunk[2:-2] if (bold or code) else chunk)
        run.bold = bold
        if code:
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def convert(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_cjk_font(doc)
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped or stripped == "---":
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            paragraph = doc.add_heading(level=level)
            add_runs(paragraph, heading.group(2))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not is_separator(lines[index]):
                    rows.append(split_row(lines[index]))
                index += 1
            width = max(len(row) for row in rows)
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for col_index in range(width):
                    text = row[col_index] if col_index < len(row) else ""
                    paragraph = cells[col_index].paragraphs[0]
                    add_runs(paragraph, text, size=TABLE_PT)
                    if row_index == 0:
                        for run in paragraph.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        if re.match(r"^[-*]\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            while index < len(lines) and (re.match(r"^[-*]\s+", lines[index].strip()) or re.match(r"^\d+\.\s+", lines[index].strip())):
                item = re.sub(r"^([-*]|\d+\.)\s+", "", lines[index].strip())
                add_runs(doc.add_paragraph(style="List Bullet"), item)
                index += 1
            continue
        if stripped.startswith(">"):
            add_runs(doc.add_paragraph(style="Intense Quote"), stripped.lstrip("> ").strip())
            index += 1
            continue
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(paragraph, stripped)
        index += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"已生成 {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--md", type=Path, required=True, help="输入的 markdown 文件")
    parser.add_argument("--out", type=Path, required=True, help="输出的 docx 路径")
    args = parser.parse_args()
    if not args.md.is_file():
        parser.error(f"输入文件不存在：{args.md}")
    convert(args.md, args.out)


if __name__ == "__main__":
    main()
